from flask import Flask, render_template, request, redirect, url_for, send_from_directory, session, flash
import sqlite3, os, json
import openpyxl
from datetime import datetime
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)
app.secret_key = "carenagens-chave-interna-2026"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BANCO = os.path.join(BASE_DIR, "carenagens.db")
PASTA_NFES = os.path.join(BASE_DIR, "nfes")
HISTORICO_FLUXO = os.path.join(BASE_DIR, "historico_fluxo.json")
BASE_ESTOQUE = os.path.join(BASE_DIR, "Base de Dados - Carenagens.xlsx")
os.makedirs(PASTA_NFES, exist_ok=True)

if not os.path.exists(HISTORICO_FLUXO):
    with open(HISTORICO_FLUXO, "w", encoding="utf-8") as arquivo_historico:
        json.dump({}, arquivo_historico, ensure_ascii=False, indent=2)

PERFIS = {
    "admin": "Administrador",
    "operador": "Operador",
    "conferente": "Conferente",
    "fornecedor": "Fornecedor"
}

PERMISSOES = {
    "inicio": "Início",
    "nova_entrega": "Nova Entrega",
    "recebimentos": "Receber Entrega",
    "pendentes": "Entradas Pendentes",
    "conferencia": "Conferência",
    "notas_fiscais": "Notas Fiscais",
    "historico": "Histórico",
    "gerenciamento": "Gerenciamento"
}

PERMISSOES_POR_PERFIL = {
    "admin": list(PERMISSOES.keys()),
    "operador": ["inicio", "recebimentos"],
    "conferente": ["inicio", "conferencia", "notas_fiscais"],
    "fornecedor": ["inicio", "nova_entrega"]
}


def carregar_catalogo_carenagens():
    """Lê a aba pecas da base XLSX e retorna referência -> descrição."""
    catalogo = {}
    if not os.path.exists(BASE_ESTOQUE):
        return catalogo
    try:
        wb = openpyxl.load_workbook(BASE_ESTOQUE, read_only=True, data_only=True)
        if "pecas" not in wb.sheetnames:
            wb.close()
            return catalogo
        ws = wb["pecas"]
        for ref, desc in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if ref is None or desc is None:
                continue
            ref, desc = str(ref).strip(), str(desc).strip()
            if ref and desc:
                catalogo[ref] = desc
        wb.close()
    except Exception:
        return {}
    return catalogo


@app.context_processor
def dados_catalogo():
    return {"catalogo_carenagens": carregar_catalogo_carenagens()}


def conectar_banco():
    c = sqlite3.connect(BANCO)
    c.row_factory = sqlite3.Row
    return c


def criar_banco():
    c = conectar_banco()
    cur = c.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS recebimentos(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            numero_recibo TEXT,
            data TEXT NOT NULL,
            responsavel TEXT NOT NULL,
            divergencia TEXT NOT NULL,
            detalhe_divergencia TEXT,
            avarias TEXT NOT NULL,
            detalhe_avarias TEXT,
            embalagem_adequada TEXT NOT NULL,
            observacoes TEXT,
            status TEXT NOT NULL,
            numero_nfe TEXT,
            responsavel_entrada TEXT,
            data_entrada TEXT,
            arquivo_nfe TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS itens(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            recebimento_id INTEGER NOT NULL,
            referencia TEXT NOT NULL,
            descricao TEXT NOT NULL,
            quantidade_entregue INTEGER NOT NULL,
            quantidade_conferida INTEGER NOT NULL,
            FOREIGN KEY(recebimento_id)
            REFERENCES recebimentos(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS usuarios(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            usuario TEXT NOT NULL UNIQUE,
            senha TEXT NOT NULL,
            nome TEXT NOT NULL,
            perfil TEXT NOT NULL DEFAULT 'operador',
            ativo INTEGER NOT NULL DEFAULT 1,
            criado_em TEXT NOT NULL,
            senha_temporaria INTEGER NOT NULL DEFAULT 0
        )
    """)

    # Registra alterações feitas pelo operador na referência/descrição.
    # O valor original do fornecedor nunca é perdido.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS alteracoes_itens(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            recebimento_id INTEGER NOT NULL,
            item_id INTEGER NOT NULL,
            referencia_original TEXT NOT NULL,
            descricao_original TEXT NOT NULL,
            referencia_atual TEXT NOT NULL,
            descricao_atual TEXT NOT NULL,
            usuario TEXT NOT NULL,
            data TEXT NOT NULL,
            visualizado_adm INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY(recebimento_id) REFERENCES recebimentos(id),
            FOREIGN KEY(item_id) REFERENCES itens(id)
        )
    """)


    # ---------------------------------------------------------
    # MIGRAÇÃO: Nº de recibo pode ficar vazio na Nova Entrega.
    # A entrega só recebe o número quando o operador a recebe.
    # ---------------------------------------------------------
    cur.execute("PRAGMA table_info(recebimentos)")
    info_rec = cur.fetchall()
    colunas_rec = [x["name"] for x in info_rec]
    numero_recibo_notnull = any(
        x["name"] == "numero_recibo" and x["notnull"] == 1
        for x in info_rec
    )

    if numero_recibo_notnull:
        cur.execute("PRAGMA foreign_keys=OFF")

        cur.execute("""
            CREATE TABLE recebimentos_migracao(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                numero_recibo TEXT,
                data TEXT NOT NULL,
                responsavel TEXT NOT NULL,
                divergencia TEXT NOT NULL,
                detalhe_divergencia TEXT,
                avarias TEXT NOT NULL,
                detalhe_avarias TEXT,
                embalagem_adequada TEXT NOT NULL,
                observacoes TEXT,
                status TEXT NOT NULL,
                numero_nfe TEXT,
                responsavel_entrada TEXT,
                data_entrada TEXT,
                arquivo_nfe TEXT,
                conferencia_resultado TEXT,
                motivo_rejeicao TEXT,
                responsavel_conferencia TEXT,
                data_conferencia TEXT
            )
        """)

        # Copia somente as colunas que existem no banco atual.
        colunas_destino = [
            "id", "numero_recibo", "data", "responsavel",
            "divergencia", "detalhe_divergencia",
            "avarias", "detalhe_avarias",
            "embalagem_adequada", "observacoes", "status",
            "numero_nfe", "responsavel_entrada", "data_entrada",
            "arquivo_nfe", "conferencia_resultado", "motivo_rejeicao",
            "responsavel_conferencia", "data_conferencia"
        ]
        colunas_existentes = set(colunas_rec)
        colunas_para_copiar = [c for c in colunas_destino if c in colunas_existentes]

        # Valores antigos que não existiam nas versões anteriores ficam NULL.
        nomes = ", ".join(colunas_para_copiar)
        cur.execute(f"""
            INSERT INTO recebimentos_migracao ({nomes})
            SELECT {nomes}
            FROM recebimentos
        """)

        cur.execute("DROP TABLE recebimentos")
        cur.execute("ALTER TABLE recebimentos_migracao RENAME TO recebimentos")
        cur.execute("PRAGMA foreign_keys=ON")

    cur.execute("PRAGMA table_info(usuarios)")
    colunas = [x["name"] for x in cur.fetchall()]

    if "senha_temporaria" not in colunas:
        cur.execute("""
            ALTER TABLE usuarios
            ADD COLUMN senha_temporaria
            INTEGER NOT NULL DEFAULT 0
        """)

    if "perfis" not in colunas:
        cur.execute("""
            ALTER TABLE usuarios
            ADD COLUMN perfis TEXT
        """)
        cur.execute("""
            UPDATE usuarios
            SET perfis = perfil
            WHERE perfis IS NULL OR perfis = ''
        """)

    cur.execute("PRAGMA table_info(recebimentos)")
    colunas_recebimentos = [x["name"] for x in cur.fetchall()]

    if "conferencia_resultado" not in colunas_recebimentos:
        cur.execute("""
            ALTER TABLE recebimentos
            ADD COLUMN conferencia_resultado TEXT
        """)

    if "motivo_rejeicao" not in colunas_recebimentos:
        cur.execute("""
            ALTER TABLE recebimentos
            ADD COLUMN motivo_rejeicao TEXT
        """)

    if "responsavel_conferencia" not in colunas_recebimentos:
        cur.execute("""
            ALTER TABLE recebimentos
            ADD COLUMN responsavel_conferencia TEXT
        """)

    if "data_conferencia" not in colunas_recebimentos:
        cur.execute("""
            ALTER TABLE recebimentos
            ADD COLUMN data_conferencia TEXT
        """)

    cur.execute("SELECT COUNT(*) total FROM usuarios")

    if cur.fetchone()["total"] == 0:
        cur.execute("""
            INSERT INTO usuarios(
                usuario,
                senha,
                nome,
                perfil,
                ativo,
                criado_em,
                senha_temporaria
            )
            VALUES(?,?,?,?,?,?,?)
        """, (
            "admin",
            generate_password_hash("admin123"),
            "Administrador",
            "admin",
            1,
            datetime.now().strftime("%d/%m/%Y %H:%M"),
            0
        ))

    c.commit()
    c.close()


def obter_perfis_usuario(usuario):
    if not usuario:
        return []

    texto = ""
    try:
        texto = usuario["perfis"] or ""
    except (KeyError, IndexError):
        texto = ""

    perfis = [p.strip() for p in texto.split(",") if p.strip() in PERFIS]

    # Compatibilidade com usuários antigos que ainda só possuem perfil.
    if not perfis and usuario["perfil"] in PERFIS:
        perfis = [usuario["perfil"]]

    return list(dict.fromkeys(perfis))


def obter_cargos(usuario):
    return obter_perfis_usuario(usuario)


def usuario_tem_perfil(usuario, perfil):
    return perfil in obter_perfis_usuario(usuario)


def permissoes_do_usuario(usuario):
    permissoes = []
    for perfil in obter_perfis_usuario(usuario):
        for permissao in PERMISSOES_POR_PERFIL.get(perfil, []):
            if permissao not in permissoes:
                permissoes.append(permissao)
    return permissoes


def obter_usuario():
    uid = session.get("usuario_id")

    if not uid:
        return None

    c = conectar_banco()

    u = c.execute("""
        SELECT *
        FROM usuarios
        WHERE id = ?
    """, (uid,)).fetchone()

    c.close()

    return u


def destino_usuario(usuario):
    if not usuario:
        return url_for("login")

    if usuario_tem_perfil(usuario, "admin"):
        return url_for("inicio")

    if usuario_tem_perfil(usuario, "operador"):
        return url_for("recebimentos")

    if usuario_tem_perfil(usuario, "conferente"):
        return url_for("conferencia")

    if usuario_tem_perfil(usuario, "fornecedor"):
        return url_for("nova_entrega")

    return url_for("login")

def possui_permissao(permissao):
    usuario = obter_usuario()

    if not usuario:
        return False

    return permissao in permissoes_do_usuario(usuario)

@app.context_processor
def contexto_global():
    """Dados globais usados pelo menu e pelos avisos do administrador."""
    usuario = obter_usuario()
    alteracoes_pendentes = 0

    if usuario and usuario_tem_perfil(usuario, "admin"):
        c = conectar_banco()
        alteracoes_pendentes = c.execute("""
            SELECT COUNT(*) AS total
            FROM alteracoes_itens
            WHERE visualizado_adm = 0
        """).fetchone()["total"]
        c.close()

    return {
        "alteracoes_pendentes": alteracoes_pendentes
    }


def login_obrigatorio(funcao):

    @wraps(funcao)
    def wrapper(*args, **kwargs):

        usuario = obter_usuario()

        if not usuario:
            return redirect(url_for("login"))

        if not usuario["ativo"]:
            session.clear()
            return redirect(url_for("login"))

        if (
            usuario["senha_temporaria"]
            and request.endpoint not in [
                "trocar_senha",
                "logout"
            ]
        ):
            return redirect(
                url_for("trocar_senha")
            )

        return funcao(*args, **kwargs)

    return wrapper


def requer_permissao(permissao):

    def decorator(funcao):

        @wraps(funcao)
        def wrapper(*args, **kwargs):

            usuario = obter_usuario()

            if not usuario:
                return redirect(
                    url_for("login")
                )

            if not possui_permissao(permissao):
                return redirect(
                    destino_usuario(usuario)
                )

            return funcao(*args, **kwargs)

        return wrapper

    return decorator


@app.context_processor
def contexto_global():

    usuario = obter_usuario()

    permissoes = permissoes_do_usuario(usuario) if usuario else []

    return {
        "usuario_logado": usuario,
        "permissoes": permissoes,
        "perfis": PERFIS,
        "permissoes_lista": PERMISSOES
    }


# =========================================================
# LOGIN
# =========================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    usuario_atual = obter_usuario()

    if usuario_atual:

        if usuario_atual["senha_temporaria"]:
            return redirect(
                url_for("trocar_senha")
            )

        return redirect(
            destino_usuario(usuario_atual)
        )

    if request.method == "POST":

        usuario = request.form.get(
            "usuario",
            ""
        ).strip().lower()

        senha = request.form.get(
            "senha",
            ""
        )

        c = conectar_banco()

        registro = c.execute("""
            SELECT *
            FROM usuarios
            WHERE usuario = ?
        """, (usuario,)).fetchone()

        c.close()

        if (
            registro
            and registro["ativo"]
            and check_password_hash(
                registro["senha"],
                senha
            )
        ):

            session.clear()

            session["usuario_id"] = registro["id"]
            session["usuario"] = registro["usuario"]
            session["nome_usuario"] = registro["nome"]

            if registro["senha_temporaria"]:
                return redirect(
                    url_for("trocar_senha")
                )

            return redirect(
                destino_usuario(registro)
            )

        return render_template(
            "login.html",
            erro="Usuário ou senha inválidos."
        )

    return render_template("login.html")


# =========================================================
# TROCAR SENHA
# =========================================================

@app.route("/trocar-senha", methods=["GET", "POST"])
@login_obrigatorio
def trocar_senha():

    usuario = obter_usuario()

    if request.method == "POST":

        nova_senha = request.form.get(
            "nova_senha",
            ""
        )

        confirmar = request.form.get(
            "confirmar_senha",
            ""
        )

        if len(nova_senha) < 6:

            return render_template(
                "trocar_senha.html",
                erro="A senha precisa ter pelo menos 6 caracteres."
            )

        if nova_senha != confirmar:

            return render_template(
                "trocar_senha.html",
                erro="As senhas não são iguais."
            )

        if nova_senha.lower() == "temp":

            return render_template(
                "trocar_senha.html",
                erro="Escolha uma senha diferente de 'temp'."
            )

        c = conectar_banco()

        c.execute("""
            UPDATE usuarios
            SET senha = ?,
                senha_temporaria = 0
            WHERE id = ?
        """, (
            generate_password_hash(nova_senha),
            usuario["id"]
        ))

        c.commit()
        c.close()

        return redirect(
            destino_usuario(
                obter_usuario()
            )
        )

    return render_template(
        "trocar_senha.html"
    )


# =========================================================
# SAIR
# =========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# =========================================================
# HISTÓRICO DE FLUXO - ARQUIVO SEPARADO DO BANCO
# =========================================================

def ler_historico_fluxo():
    try:
        with open(HISTORICO_FLUXO, "r", encoding="utf-8") as arquivo:
            dados = json.load(arquivo)
            return dados if isinstance(dados, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def registrar_evento_fluxo(recebimento_id, etapa, usuario, data=None, detalhe=""):
    dados = ler_historico_fluxo()
    chave = str(recebimento_id)
    if chave not in dados:
        dados[chave] = []
    dados[chave].append({
        "etapa": etapa,
        "usuario": usuario or "Usuário não identificado",
        "data": data or datetime.now().strftime("%d/%m/%Y %H:%M"),
        "detalhe": detalhe or ""
    })
    temporario = HISTORICO_FLUXO + ".tmp"
    try:
        with open(temporario, "w", encoding="utf-8") as arquivo:
            json.dump(dados, arquivo, ensure_ascii=False, indent=2)
        os.replace(temporario, HISTORICO_FLUXO)
    except OSError:
        try:
            if os.path.exists(temporario): os.remove(temporario)
        except OSError:
            pass


def eventos_fluxo(recebimento):
    dados = ler_historico_fluxo()
    eventos = list(dados.get(str(recebimento["id"]), []))
    if not eventos:
        eventos.append({"etapa":"Recebimento criado","usuario":recebimento["responsavel"],"data":recebimento["data"],"detalhe":""})
        # Para registros antigos sem evento de NF-e, não inventa usuário.
        # Registros novos terão o usuário salvo pelo registrar_evento_fluxo().
        if recebimento["numero_nfe"]:
            eventos.append({
                "etapa": "NF-e anexada",
                "usuario": "",
                "data": "",
                "detalhe": f"NF-e nº {recebimento['numero_nfe']}"
            })
        if recebimento["responsavel_conferencia"]:
            eventos.append({"etapa":"Conferência aprovada" if recebimento["conferencia_resultado"] == "aprovado" else "Conferência rejeitada","usuario":recebimento["responsavel_conferencia"],"data":recebimento["data_conferencia"] or "","detalhe":recebimento["motivo_rejeicao"] or ""})
        if recebimento["responsavel_entrada"]:
            eventos.append({"etapa":"Entrada realizada","usuario":recebimento["responsavel_entrada"],"data":recebimento["data_entrada"] or "","detalhe":""})
    return eventos


# =========================================================
# INÍCIO - PAINEL DE ACOMPANHAMENTO
# =========================================================

def calcular_status_inicio(recebimento):
    """
    Status amigável mostrado no painel inicial.
    Não altera o banco: apenas interpreta os dados existentes.
    """
    if recebimento["status"] == "Aguardando recebimento":
        return "Aguardando recebimento"
    if not recebimento["numero_nfe"]:
        return "Aguardando NF-e"
    if recebimento["status"] == "Aguardando conferência":
        return "Aguardando conferência"
    if recebimento["status"] == "Aguardando entrada":
        return "Aguardando entrada"
    if recebimento["status"] == "Finalizado":
        return "Finalizado"
    return recebimento["status"]


def calcular_ultima_etapa(recebimento):
    if recebimento["responsavel_entrada"]:
        return f"Entrada realizada por {recebimento['responsavel_entrada']}"
    if recebimento["responsavel_conferencia"]:
        resultado = recebimento["conferencia_resultado"] or "registrada"
        if resultado == "aprovado":
            return f"Conferência aprovada por {recebimento['responsavel_conferencia']}"
        if resultado == "rejeitado":
            return f"Conferência rejeitada por {recebimento['responsavel_conferencia']}"
        return f"Conferência registrada por {recebimento['responsavel_conferencia']}"
    if recebimento["numero_nfe"]:
        return f"NF-e nº {recebimento['numero_nfe']} anexada"
    return f"Recebimento criado por {recebimento['responsavel']}"


def preparar_registro_inicio(recebimento):
    registro = dict(recebimento)
    registro["status_inicio"] = calcular_status_inicio(recebimento)
    registro["ultima_etapa"] = calcular_ultima_etapa(recebimento)
    registro.setdefault("responsavel_nfe", None)
    registro.setdefault("data_nfe", None)
    return registro


def carregar_registros_inicio():
    c = conectar_banco()
    registros = c.execute("""
        SELECT *
        FROM recebimentos
        ORDER BY id DESC
    """).fetchall()
    c.close()
    return [preparar_registro_inicio(r) for r in registros]


@app.route("/")
@login_obrigatorio
@requer_permissao("inicio")
def inicio():
    registros = carregar_registros_inicio()
    contagens = {
        "Aguardando NF-e": 0,
        "Aguardando conferência": 0,
        "Aguardando entrada": 0,
        "Finalizado": 0
    }
    for registro in registros:
        status = registro["status_inicio"]
        if status in contagens:
            contagens[status] += 1
    return render_template(
        "index.html",
        pagina="inicio",
        aguardando_nfe=contagens["Aguardando NF-e"],
        aguardando_conferencia=contagens["Aguardando conferência"],
        aguardando_entrada=contagens["Aguardando entrada"],
        finalizados=contagens["Finalizado"],
        registros_inicio=registros,
        total_pedidos=len(registros)
    )


@app.route("/api/inicio")
@login_obrigatorio
@requer_permissao("inicio")
def api_inicio():
    registros = carregar_registros_inicio()
    contagens = {
        "Aguardando NF-e": 0,
        "Aguardando conferência": 0,
        "Aguardando entrada": 0,
        "Finalizado": 0
    }
    dados = []
    for registro in registros:
        status = registro["status_inicio"]
        if status in contagens:
            contagens[status] += 1
        dados.append({
            "id": registro["id"],
            "numero_recibo": registro["numero_recibo"],
            "data": registro["data"],
            "responsavel": registro["responsavel"],
            "numero_nfe": registro["numero_nfe"],
            "status_inicio": status,
            "ultima_etapa": registro["ultima_etapa"]
        })
    return {
        "contagens": contagens,
        "registros": dados,
        "total": len(dados),
        "atualizado_em": datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    }


@app.route("/pedido/<int:recebimento_id>")
@login_obrigatorio
@requer_permissao("inicio")
def detalhe_pedido(recebimento_id):
    c = conectar_banco()
    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()
    if not recebimento:
        c.close()
        flash("Recebimento não encontrado.", "erro")
        return redirect(url_for("inicio"))
    itens_pedido = c.execute("""
        SELECT *
        FROM itens
        WHERE recebimento_id = ?
        ORDER BY id
    """, (recebimento_id,)).fetchall()
    colunas = [x["name"] for x in c.execute("PRAGMA table_info(recebimentos)").fetchall()]
    recebimento_dict = dict(recebimento)
    recebimento_dict["responsavel_nfe"] = recebimento["responsavel_nfe"] if "responsavel_nfe" in colunas else None
    recebimento_dict["data_nfe"] = recebimento["data_nfe"] if "data_nfe" in colunas else None
    c.close()
    status_inicio = calcular_status_inicio(recebimento)
    return render_template(
        "index.html",
        pagina="pedido-detalhe",
        recebimento=recebimento_dict,
        itens_pedido=itens_pedido,
        status_inicio=status_inicio,
        eventos=eventos_fluxo(recebimento)
    )


@app.route("/api/pedido/<int:recebimento_id>")
@login_obrigatorio
@requer_permissao("inicio")
def api_detalhe_pedido(recebimento_id):
    c = conectar_banco()
    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()
    if not recebimento:
        c.close()
        return {"erro": "Recebimento não encontrado."}, 404
    itens = c.execute("""
        SELECT *
        FROM itens
        WHERE recebimento_id = ?
        ORDER BY id
    """, (recebimento_id,)).fetchall()
    c.close()
    return {
        "id": recebimento["id"],
        "numero_recibo": recebimento["numero_recibo"],
        "data": recebimento["data"],
        "responsavel": recebimento["responsavel"],
        "numero_nfe": recebimento["numero_nfe"],
        "arquivo_nfe": recebimento["arquivo_nfe"],
        "status": calcular_status_inicio(recebimento),
        "conferencia_resultado": recebimento["conferencia_resultado"],
        "motivo_rejeicao": recebimento["motivo_rejeicao"],
        "responsavel_conferencia": recebimento["responsavel_conferencia"],
        "data_conferencia": recebimento["data_conferencia"],
        "responsavel_entrada": recebimento["responsavel_entrada"],
        "data_entrada": recebimento["data_entrada"],
        "eventos": eventos_fluxo(recebimento),
        "atualizado_em": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "itens": [dict(item) for item in itens]
    }



# =========================================================
# NOVA ENTREGA - FORNECEDOR
# =========================================================

@app.route("/nova-entrega")
@login_obrigatorio
@requer_permissao("nova_entrega")
def nova_entrega():
    return render_template(
        "index.html",
        pagina="nova-entrega",
        data_atual=datetime.now().strftime("%Y-%m-%d")
    )


@app.route("/salvar-nova-entrega", methods=["POST"])
@login_obrigatorio
@requer_permissao("nova_entrega")
def salvar_nova_entrega():
    usuario = obter_usuario()
    if not usuario:
        return redirect(url_for("login"))

    data = request.form.get("data", "").strip()
    referencias = request.form.getlist("referencia[]")
    descricoes = request.form.getlist("descricao[]")
    entregues = request.form.getlist("quantidade_entregue[]")

    if not data:
        flash("Informe a data da entrega.", "erro")
        return redirect(url_for("nova_entrega"))

    catalogo = carregar_catalogo_carenagens()
    itens = []
    for i, referencia in enumerate(referencias):
        referencia = referencia.strip()
        try:
            quantidade = int(entregues[i])
        except (ValueError, TypeError):
            quantidade = 0

        if not referencia:
            continue
        if referencia not in catalogo:
            flash(f"Referência não encontrada na base de carenagens: {referencia}", "erro")
            return redirect(url_for("nova_entrega"))
        if quantidade < 0:
            flash("Informe uma quantidade entregue válida.", "erro")
            return redirect(url_for("nova_entrega"))
        itens.append((referencia, catalogo[referencia], quantidade))

    if not itens:
        flash("Adicione pelo menos uma peça.", "erro")
        return redirect(url_for("nova_entrega"))

    agora = datetime.now().strftime("%d/%m/%Y %H:%M")
    c = conectar_banco()
    cur = c.cursor()

    cur.execute("""
        INSERT INTO recebimentos(
            numero_recibo, data, responsavel,
            divergencia, detalhe_divergencia,
            avarias, detalhe_avarias,
            embalagem_adequada, observacoes, status
        )
        VALUES(?,?,?,?,?,?,?,?,?,?)
    """, (
        None,
        data, usuario["nome"], "", "", "", "", "", "",
        "Aguardando recebimento"
    ))

    recebimento_id = cur.lastrowid

    for referencia, descricao, quantidade in itens:
        cur.execute("""
            INSERT INTO itens(
                recebimento_id, referencia, descricao,
                quantidade_entregue, quantidade_conferida
            )
            VALUES(?,?,?,?,?)
        """, (recebimento_id, referencia, descricao, quantidade, 0))

    c.commit()
    c.close()

    registrar_evento_fluxo(
        recebimento_id,
        "Nova entrega registrada",
        usuario["nome"],
        agora,
        "Entrega criada pelo fornecedor"
    )

    flash("Entrega registrada com sucesso.", "sucesso")
    return redirect(url_for("nova_entrega"))


# =========================================================
# RECEBIMENTO / CONFERÊNCIA DO OPERADOR
# =========================================================

@app.route("/recebimento-operador/<int:recebimento_id>")
@login_obrigatorio
@requer_permissao("recebimentos")
def recebimento_operador(recebimento_id):
    usuario = obter_usuario()
    if not usuario or not (usuario_tem_perfil(usuario, "operador") or usuario_tem_perfil(usuario, "admin")):
        return redirect(destino_usuario(usuario))

    c = conectar_banco()
    recebimento = c.execute(
        "SELECT * FROM recebimentos WHERE id = ?",
        (recebimento_id,)
    ).fetchone()
    itens = c.execute(
        "SELECT * FROM itens WHERE recebimento_id = ? ORDER BY id",
        (recebimento_id,)
    ).fetchall()
    c.close()

    if not recebimento:
        flash("Entrega não encontrada.", "erro")
        return redirect(url_for("recebimentos"))

    if recebimento["status"] != "Aguardando recebimento":
        flash("Esta entrega não está aguardando recebimento pelo operador.", "erro")
        return redirect(url_for("recebimentos"))

    return render_template(
        "index.html",
        pagina="recebimento-operador",
        recebimento=recebimento,
        itens=itens
    )


@app.route("/salvar-conferencia-operador/<int:recebimento_id>", methods=["POST"])
@login_obrigatorio
@requer_permissao("recebimentos")
def salvar_conferencia_operador(recebimento_id):
    usuario = obter_usuario()
    if not usuario or not (usuario_tem_perfil(usuario, "operador") or usuario_tem_perfil(usuario, "admin")):
        return redirect(destino_usuario(usuario))

    c = conectar_banco()
    recebimento = c.execute(
        "SELECT * FROM recebimentos WHERE id = ?",
        (recebimento_id,)
    ).fetchone()
    itens_db = c.execute(
        "SELECT * FROM itens WHERE recebimento_id = ? ORDER BY id",
        (recebimento_id,)
    ).fetchall()

    if not recebimento:
        c.close()
        flash("Entrega não encontrada.", "erro")
        return redirect(url_for("recebimentos"))

    if recebimento["status"] != "Aguardando recebimento":
        c.close()
        flash("Esta entrega não está aguardando recebimento pelo operador.", "erro")
        return redirect(url_for("recebimentos"))

    numero_recibo = request.form.get("numero_recibo", "").strip()
    referencias = request.form.getlist("referencia[]")
    descricoes = request.form.getlist("descricao[]")
    conferidas = request.form.getlist("quantidade_conferida[]")
    divergencia = request.form.get("divergencia", "Não").strip()
    avarias = request.form.get("avarias", "Não").strip()
    embalagem_adequada = request.form.get("embalagem_adequada", "Sim").strip()
    observacoes = request.form.get("observacoes", "").strip()

    if not numero_recibo:
        c.close()
        flash("Informe o Nº do recibo.", "erro")
        return redirect(url_for("recebimento_operador", recebimento_id=recebimento_id))

    if len(referencias) != len(itens_db) or len(descricoes) != len(itens_db) or len(conferidas) != len(itens_db):
        c.close()
        flash("Não foi possível validar os itens da entrega.", "erro")
        return redirect(url_for("recebimento_operador", recebimento_id=recebimento_id))

    alteracoes = []
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")

    catalogo = carregar_catalogo_carenagens()
    for item, referencia, descricao, quantidade in zip(itens_db, referencias, descricoes, conferidas):
        referencia = referencia.strip()
        if referencia not in catalogo:
            c.close()
            flash(f"Referência não encontrada na base de carenagens: {referencia}", "erro")
            return redirect(url_for("recebimento_operador", recebimento_id=recebimento_id))
        descricao = catalogo[referencia]
        try:
            quantidade_conferida = int(quantidade)
        except (ValueError, TypeError):
            c.close()
            flash("Informe uma quantidade conferida válida.", "erro")
            return redirect(url_for("recebimento_operador", recebimento_id=recebimento_id))

        if not referencia or not descricao or quantidade_conferida < 0:
            c.close()
            flash("Referência, descrição e quantidade conferida são obrigatórias.", "erro")
            return redirect(url_for("recebimento_operador", recebimento_id=recebimento_id))

        mudou_referencia = referencia != item["referencia"]
        mudou_descricao = descricao != item["descricao"]

        if mudou_referencia or mudou_descricao:
            alteracoes.append({
                "item_id": item["id"],
                "referencia_original": item["referencia"],
                "descricao_original": item["descricao"],
                "referencia_atual": referencia,
                "descricao_atual": descricao
            })

        c.execute("""
            UPDATE itens
            SET referencia = ?, descricao = ?, quantidade_conferida = ?
            WHERE id = ?
        """, (referencia, descricao, quantidade_conferida, item["id"]))

    c.execute("""
        UPDATE recebimentos
        SET
            numero_recibo = ?,
            divergencia = ?,
            avarias = ?,
            embalagem_adequada = ?,
            observacoes = ?,
            status = 'Aguardando entrada'
        WHERE id = ?
    """, (
        numero_recibo,
        divergencia,
        avarias,
        embalagem_adequada,
        observacoes,
        recebimento_id
    ))

    for alteracao in alteracoes:
        c.execute("""
            INSERT INTO alteracoes_itens(
                recebimento_id, item_id,
                referencia_original, descricao_original,
                referencia_atual, descricao_atual,
                usuario, data, visualizado_adm
            ) VALUES(?,?,?,?,?,?,?,?,0)
        """, (
            recebimento_id,
            alteracao["item_id"],
            alteracao["referencia_original"],
            alteracao["descricao_original"],
            alteracao["referencia_atual"],
            alteracao["descricao_atual"],
            usuario["nome"],
            agora
        ))

    c.commit()
    c.close()

    detalhe = "Conferência do operador concluída."
    if alteracoes:
        partes = []
        for a in alteracoes:
            partes.append(
                f"Ref.: {a['referencia_original']} → {a['referencia_atual']} | "
                f"Desc.: {a['descricao_original']} → {a['descricao_atual']}"
            )
        detalhe += " Alterações informadas ao administrador: " + " || ".join(partes)

    registrar_evento_fluxo(
        recebimento_id,
        "Conferência do operador",
        usuario["nome"],
        agora,
        detalhe
    )

    if alteracoes:
        flash("Conferência registrada. O administrador foi sinalizado sobre as alterações de referência/descrição.", "sucesso")
    else:
        flash("Conferência do operador registrada. Enviada para o administrador.", "sucesso")

    return redirect(url_for("recebimentos"))


# =========================================================
# NOVO RECEBIMENTO
# =========================================================
# NOVO RECEBIMENTO
# =========================================================

@app.route("/recebimentos")
@login_obrigatorio
@requer_permissao("recebimentos")
def recebimentos():

    usuario = obter_usuario()

    if usuario and (usuario_tem_perfil(usuario, "operador") or usuario_tem_perfil(usuario, "admin")):
        c = conectar_banco()
        registros = c.execute("""
            SELECT *
            FROM recebimentos
            WHERE status = 'Aguardando recebimento'
            ORDER BY id DESC
        """).fetchall()
        c.close()

        # Usa o mesmo layout principal do sistema para manter o menu lateral
        # visível também na fila de recebimentos do operador.
        return render_template(
            "index.html",
            pagina="recebimentos",
            registros=registros
        )

    return render_template(
        "index.html",
        pagina="recebimentos",
        data_atual=datetime.now().strftime("%Y-%m-%d")
    )


@app.route(
    "/salvar-recebimento",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("recebimentos")
def salvar_recebimento():

    f = request.form

    numero = f.get(
        "numero_recibo",
        ""
    ).strip()

    data = f.get(
        "data",
        ""
    ).strip()

    usuario = obter_usuario()

    if not usuario:
        return redirect(url_for("login"))

    # O responsável da solicitação é automaticamente o usuário logado.
    # Nada é preenchido manualmente e nenhum usuário é inventado.
    responsavel = usuario["nome"]

    if not numero or not data:
        return """
        <script>
        alert("Preencha Nº do recibo e data.");
        history.back();
        </script>
        """

    referencias = f.getlist(
        "referencia[]"
    )

    descricoes = f.getlist(
        "descricao[]"
    )

    entregues = f.getlist(
        "quantidade_entregue[]"
    )

    conferidas = f.getlist(
        "quantidade_conferida[]"
    )

    itens = []

    for i, referencia in enumerate(
        referencias
    ):

        descricao = descricoes[i].strip()

        if not referencia.strip() and not descricao:
            continue

        if not referencia.strip() or not descricao:

            return """
            <script>
            alert("Preencha referência e descrição de todos os itens.");
            history.back();
            </script>
            """

        try:
            quantidade_entregue = int(
                entregues[i]
            )
        except:
            quantidade_entregue = 0

        try:
            quantidade_conferida = int(
                conferidas[i]
            )
        except:
            quantidade_conferida = 0

        itens.append((
            referencia.strip(),
            descricao,
            quantidade_entregue,
            quantidade_conferida
        ))

    if not itens:

        return """
        <script>
        alert("Adicione pelo menos um item.");
        history.back();
        </script>
        """

    c = conectar_banco()

    cur = c.cursor()

    cur.execute("""
        INSERT INTO recebimentos(
            numero_recibo,
            data,
            responsavel,
            divergencia,
            detalhe_divergencia,
            avarias,
            detalhe_avarias,
            embalagem_adequada,
            observacoes,
            status
        )
        VALUES(?,?,?,?,?,?,?,?,?,?)
    """, (
        numero,
        data,
        responsavel,
        f.get("divergencia", "Não"),
        f.get(
            "detalhe_divergencia",
            ""
        ).strip(),
        f.get("avarias", "Não"),
        f.get(
            "detalhe_avarias",
            ""
        ).strip(),
        f.get(
            "embalagem_adequada",
            "Sim"
        ),
        f.get(
            "observacoes",
            ""
        ).strip(),
        "Aguardando entrada"
    ))

    recebimento_id = cur.lastrowid

    for item in itens:

        cur.execute("""
            INSERT INTO itens(
                recebimento_id,
                referencia,
                descricao,
                quantidade_entregue,
                quantidade_conferida
            )
            VALUES(?,?,?,?,?)
        """, (
            recebimento_id,
            item[0],
            item[1],
            item[2],
            item[3]
        ))

    c.commit()
    c.close()

    registrar_evento_fluxo(
        recebimento_id,
        "Solicitação registrada",
        usuario["nome"],
        data,
        ""
    )

    return redirect(
        url_for("pendentes")
    )


# =========================================================
# ENTRADAS PENDENTES
# =========================================================

@app.route("/pendentes")
@login_obrigatorio
@requer_permissao("pendentes")
def pendentes():
    c = conectar_banco()
    registros = c.execute("""
        SELECT r.*,
               EXISTS(SELECT 1 FROM alteracoes_itens a
                      WHERE a.recebimento_id = r.id
                        AND a.visualizado_adm = 0) AS tem_alteracao_pendente
        FROM recebimentos r
        WHERE r.status = 'Aguardando entrada'
           OR (r.status = 'Aguardando conferência'
               AND (r.numero_nfe IS NULL OR r.numero_nfe = ''))
        ORDER BY r.id DESC
    """).fetchall()
    c.close()

    return render_template(
        "index.html",
        pagina="pendentes",
        registros=registros
    )


@app.route("/recebimento/<int:recebimento_id>")
@login_obrigatorio
@requer_permissao("pendentes")
def abrir_recebimento(recebimento_id):
    c = conectar_banco()

    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()

    itens = c.execute("""
        SELECT *
        FROM itens
        WHERE recebimento_id = ?
        ORDER BY id
    """, (recebimento_id,)).fetchall()

    alteracoes = c.execute("""
        SELECT *
        FROM alteracoes_itens
        WHERE recebimento_id = ?
        ORDER BY id
    """, (recebimento_id,)).fetchall()

    if not recebimento:
        c.close()
        return """
        <script>
        alert("Recebimento não encontrado.");
        window.location.href="/pendentes";
        </script>
        """

    # Ao abrir a entrada, o administrador tomou ciência da alteração.
    c.execute("""
        UPDATE alteracoes_itens
        SET visualizado_adm = 1
        WHERE recebimento_id = ?
    """, (recebimento_id,))
    c.commit()
    c.close()

    return render_template(
        "index.html",
        pagina="entrada",
        recebimento=recebimento,
        itens=itens,
        alteracoes=alteracoes
    )


# =========================================================
# CONFIRMAR ENTRADA / NF
# =========================================================

@app.route(
    "/confirmar-entrada/<int:recebimento_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("pendentes")
def confirmar_entrada(recebimento_id):

    c = conectar_banco()

    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()

    if not recebimento:
        c.close()
        flash("Recebimento não encontrado.", "erro")
        return redirect(url_for("pendentes"))

    if recebimento["status"] != "Aguardando entrada":
        c.close()
        flash("Este recebimento não está em Entradas Pendentes.", "erro")
        return redirect(url_for("pendentes"))

    if recebimento["conferencia_resultado"] != "aprovado":
        c.close()
        flash(
            "A entrada só pode ser confirmada depois da aprovação do conferente.",
            "erro"
        )
        return redirect(url_for("pendentes"))

    if not recebimento["numero_nfe"] or not recebimento["arquivo_nfe"]:
        c.close()
        flash(
            "Este recebimento precisa ter uma NF-e anexada antes da confirmação.",
            "erro"
        )
        return redirect(url_for("pendentes"))

    c.execute("""
        UPDATE recebimentos
        SET
            responsavel_entrada = ?,
            data_entrada = ?,
            status = 'Finalizado'
        WHERE id = ?
    """, (
        obter_usuario()["nome"],
        datetime.now().strftime("%d/%m/%Y %H:%M"),
        recebimento_id
    ))

    c.commit()
    c.close()

    registrar_evento_fluxo(recebimento_id, "Entrada realizada", obter_usuario()["nome"], datetime.now().strftime("%d/%m/%Y %H:%M"))

    flash(
        "Entrada confirmada com sucesso. Recebimento finalizado.",
        "sucesso"
    )

    return redirect(url_for("pendentes"))


@app.route(
    "/nfe/<path:nome_arquivo>"
)
@login_obrigatorio
@requer_permissao("notas_fiscais")
def abrir_nfe(nome_arquivo):

    return send_from_directory(
        PASTA_NFES,
        nome_arquivo
    )


# =========================================================
# CONFERÊNCIA
# =========================================================

@app.route("/conferencia")
@login_obrigatorio
@requer_permissao("conferencia")
def conferencia():

    c = conectar_banco()

    registros = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE status = 'Aguardando conferência'
          AND numero_nfe IS NOT NULL
          AND numero_nfe != ''
        ORDER BY id DESC
    """).fetchall()

    c.close()

    return render_template(
        "index.html",
        pagina="conferencia",
        registros=registros
    )


@app.route(
    "/conferencia/<int:recebimento_id>"
)
@login_obrigatorio
@requer_permissao("conferencia")
def abrir_conferencia(recebimento_id):

    c = conectar_banco()

    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()

    # Busca as peças deste recebimento para a tela de conferência.
    itens = c.execute("""
        SELECT *
        FROM itens
        WHERE recebimento_id = ?
        ORDER BY id
    """, (recebimento_id,)).fetchall()

    c.close()

    if not recebimento:
        flash("Recebimento não encontrado.", "erro")
        return redirect(url_for("conferencia"))

    if recebimento["status"] != "Aguardando conferência":
        flash(
            "Este recebimento não está aguardando conferência.",
            "erro"
        )
        return redirect(url_for("conferencia"))

    return render_template(
        "index.html",
        pagina="conferencia-detalhe",
        recebimento=recebimento,
        itens=itens
    )


@app.route(
    "/aprovar-conferencia/<int:recebimento_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("conferencia")
def aprovar_conferencia(recebimento_id):
    usuario = obter_usuario()
    if not usuario:
        return redirect(url_for("login"))

    c = conectar_banco()
    recebimento = c.execute("""
        SELECT * FROM recebimentos WHERE id = ?
    """, (recebimento_id,)).fetchone()

    if not recebimento:
        c.close()
        flash("Recebimento não encontrado.", "erro")
        return redirect(url_for("conferencia"))

    if recebimento["status"] != "Aguardando conferência":
        c.close()
        flash("Este recebimento não está aguardando conferência.", "erro")
        return redirect(url_for("conferencia"))

    if not recebimento["numero_nfe"] or not recebimento["arquivo_nfe"]:
        c.close()
        flash("Não é possível aprovar sem uma NF-e anexada.", "erro")
        return redirect(url_for("conferencia"))

    confirmacao = request.form.get("confirmacao_nfe")
    if confirmacao != "sim":
        c.close()
        flash("Confirme que a NF-e foi comparada com o recebimento antes de aprovar.", "erro")
        return redirect(url_for("abrir_conferencia", recebimento_id=recebimento_id))

    agora = datetime.now().strftime("%d/%m/%Y %H:%M")

    # O conferente somente valida a NF-e.
    # Divergência, avaria, embalagem, observações e quantidades são
    # responsabilidade do operador e não são alteradas nesta etapa.
    c.execute("""
        UPDATE recebimentos
        SET
            status = 'Aguardando entrada',
            conferencia_resultado = 'aprovado',
            motivo_rejeicao = NULL,
            responsavel_conferencia = ?,
            data_conferencia = ?
        WHERE id = ?
    """, (
        usuario["nome"],
        agora,
        recebimento_id
    ))

    c.commit()
    c.close()

    registrar_evento_fluxo(
        recebimento_id,
        "Conferência da NF-e aprovada",
        usuario["nome"],
        agora,
        "Conferente comparou a NF-e com o recebimento e aprovou. Nenhum dado da conferência física foi alterado."
    )

    flash(
        "NF-e conferida e aprovada. O recebimento voltou para o administrador confirmar a entrada.",
        "sucesso"
    )

    return redirect(url_for("conferencia"))


@app.route(
    "/rejeitar-conferencia/<int:recebimento_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("conferencia")
def rejeitar_conferencia(recebimento_id):
    usuario = obter_usuario()
    if not usuario:
        return redirect(url_for("login"))

    motivo = request.form.get("motivo_rejeicao", "").strip()
    if not motivo:
        flash("Informe o motivo da rejeição da NF-e.", "erro")
        return redirect(url_for("abrir_conferencia", recebimento_id=recebimento_id))

    c = conectar_banco()
    recebimento = c.execute("""
        SELECT * FROM recebimentos WHERE id = ?
    """, (recebimento_id,)).fetchone()

    if not recebimento:
        c.close()
        flash("Recebimento não encontrado.", "erro")
        return redirect(url_for("conferencia"))

    if recebimento["status"] != "Aguardando conferência":
        c.close()
        flash("Este recebimento não está aguardando conferência.", "erro")
        return redirect(url_for("conferencia"))

    if recebimento["arquivo_nfe"]:
        caminho = os.path.join(PASTA_NFES, recebimento["arquivo_nfe"])
        try:
            if os.path.isfile(caminho):
                os.remove(caminho)
        except OSError:
            pass

    agora = datetime.now().strftime("%d/%m/%Y %H:%M")
    c.execute("""
        UPDATE recebimentos
        SET
            status = 'Aguardando entrada',
            numero_nfe = NULL,
            arquivo_nfe = NULL,
            conferencia_resultado = 'rejeitado',
            motivo_rejeicao = ?,
            responsavel_conferencia = ?,
            data_conferencia = ?
        WHERE id = ?
    """, (motivo, usuario["nome"], agora, recebimento_id))

    c.commit()
    c.close()

    registrar_evento_fluxo(
        recebimento_id,
        "Conferência da NF-e rejeitada",
        usuario["nome"],
        agora,
        motivo
    )

    flash(
        "NF-e rejeitada. O recebimento voltou para o administrador.",
        "sucesso"
    )
    return redirect(url_for("conferencia"))


# =========================================================
# NOTAS FISCAIS
# =========================================================

@app.route("/notas-fiscais")
@login_obrigatorio
@requer_permissao("notas_fiscais")
def notas_fiscais():

    c = conectar_banco()

    registros = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE numero_nfe IS NOT NULL
        AND numero_nfe != ''
        ORDER BY id DESC
    """).fetchall()

    c.close()

    return render_template(
        "index.html",
        pagina="notas-fiscais",
        registros=registros
    )


@app.route(
    "/anexar-nfe/<int:recebimento_id>",
    methods=["POST"]
)
@login_obrigatorio
def anexar_nfe(recebimento_id):

    usuario = obter_usuario()

    if not usuario or not usuario_tem_perfil(usuario, "admin"):
        return redirect(destino_usuario(usuario))

    numero_nfe = request.form.get("numero_nfe", "").strip()
    arquivo = request.files.get("arquivo_nfe")

    if not numero_nfe:
        flash("Informe o número da NF-e.", "erro")
        return redirect(url_for("pendentes"))

    if not arquivo or not arquivo.filename:
        flash("Selecione o arquivo da NF-e.", "erro")
        return redirect(url_for("pendentes"))

    c = conectar_banco()

    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()

    if not recebimento:
        c.close()
        flash("Recebimento não encontrado.", "erro")
        return redirect(url_for("pendentes"))

    if recebimento["status"] not in ["Aguardando entrada", "Aguardando conferência"]:
        c.close()
        flash("A NF-e não pode ser anexada nesta etapa.", "erro")
        return redirect(url_for("pendentes"))

    if recebimento["status"] == "Aguardando conferência":
        eventos = ler_historico_fluxo().get(str(recebimento_id), [])
        if not any(e.get("etapa") == "Conferência do operador" for e in eventos):
            c.close()
            flash("A entrega ainda precisa ser conferida pelo operador.", "erro")
            return redirect(url_for("pendentes"))

    extensao = os.path.splitext(arquivo.filename)[1].lower()

    if extensao not in [".pdf", ".jpg", ".jpeg", ".png"]:
        c.close()
        flash("Formato de arquivo não permitido.", "erro")
        return redirect(url_for("pendentes"))

    if recebimento["arquivo_nfe"]:
        caminho_antigo = os.path.join(
            PASTA_NFES,
            recebimento["arquivo_nfe"]
        )
        try:
            if os.path.isfile(caminho_antigo):
                os.remove(caminho_antigo)
        except OSError:
            pass

    nome_arquivo = (
        f"recebimento_{recebimento_id}_"
        f"nfe_{numero_nfe}{extensao}"
    )

    arquivo.save(
        os.path.join(PASTA_NFES, nome_arquivo)
    )

    c.execute("""
        UPDATE recebimentos
        SET
            numero_nfe = ?,
            arquivo_nfe = ?,
            status = ?,
            conferencia_resultado = NULL,
            motivo_rejeicao = NULL
        WHERE id = ?
    """, (
        numero_nfe,
        nome_arquivo,
        "Aguardando conferência",
        recebimento_id
    ))

    c.commit()
    c.close()

    registrar_evento_fluxo(recebimento_id, "NF-e anexada", usuario["nome"], datetime.now().strftime("%d/%m/%Y %H:%M"), f"NF-e nº {numero_nfe}")

    flash(
        "NF-e anexada. O recebimento foi enviado para Conferência.",
        "sucesso"
    )

    return redirect(url_for("pendentes"))


# =========================================================
# HISTÓRICO
# =========================================================

@app.route("/historico")
@login_obrigatorio
@requer_permissao("historico")
def historico():

    c = conectar_banco()

    registros = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE status = 'Finalizado' 
        ORDER BY id DESC
    """).fetchall()

    c.close()

    return render_template(
        "index.html",
        pagina="historico",
        registros=registros
    )


# =========================================================
# EXCLUIR REGISTRO DO HISTÓRICO
# =========================================================

@app.route("/excluir-historico/<int:recebimento_id>", methods=["POST"])
@login_obrigatorio
@requer_permissao("historico")
def excluir_historico(recebimento_id):

    c = conectar_banco()

    registro = c.execute("""
        SELECT id, status, arquivo_nfe
        FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,)).fetchone()

    if not registro:
        c.close()
        flash("Registro não encontrado.", "erro")
        return redirect(url_for("historico"))

    if registro["status"] != "Finalizado":
        c.close()
        flash("Somente registros finalizados podem ser excluídos do histórico.", "erro")
        return redirect(url_for("historico"))

    # Remove o arquivo da NF-e, se existir.
    if registro["arquivo_nfe"]:
        caminho = os.path.join(PASTA_NFES, registro["arquivo_nfe"])
        try:
            if os.path.isfile(caminho):
                os.remove(caminho)
        except OSError:
            pass

    # Remove os itens vinculados e o recebimento.
    c.execute("""
        DELETE FROM itens
        WHERE recebimento_id = ?
    """, (recebimento_id,))

    c.execute("""
        DELETE FROM recebimentos
        WHERE id = ?
    """, (recebimento_id,))

    c.commit()
    c.close()

    flash("Registro excluído do histórico.", "sucesso")
    return redirect(url_for("historico"))



# =========================================================
# GERENCIAMENTO
# =========================================================

@app.route("/gerenciamento")
@login_obrigatorio
@requer_permissao("gerenciamento")
def gerenciamento():

    c = conectar_banco()

    usuarios = c.execute("""
        SELECT *
        FROM usuarios
        ORDER BY nome
    """).fetchall()

    c.close()

    usuario_perfis_map = {u["id"]: obter_perfis_usuario(u) for u in usuarios}

    return render_template(
        "gerenciamento.html",
        pagina="gerenciamento",
        usuarios=usuarios,
        usuario_perfis_map=usuario_perfis_map
    )


@app.route(
    "/criar-usuario",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def criar_usuario():

    usuario = request.form.get("usuario", "").strip().lower()
    nome = request.form.get("nome", "").strip()
    senha = request.form.get("senha", "")

    perfis_selecionados = [
        p for p in request.form.getlist("perfis")
        if p in PERFIS
    ]

    if not perfis_selecionados:
        perfil_legado = request.form.get("perfil", "operador")
        if perfil_legado in PERFIS:
            perfis_selecionados = [perfil_legado]
        else:
            perfis_selecionados = ["operador"]

    # Administrador primeiro para manter compatibilidade com o campo perfil antigo.
    if "admin" in perfis_selecionados:
        perfis_selecionados = ["admin"] + [p for p in perfis_selecionados if p != "admin"]

    perfil_principal = perfis_selecionados[0]
    perfis_texto = ",".join(perfis_selecionados)

    if not usuario or not nome or not senha:
        flash("Preencha todos os campos.", "erro")
        return redirect(url_for("gerenciamento"))

    if len(senha) < 6:
        flash("A senha precisa ter pelo menos 6 caracteres.", "erro")
        return redirect(url_for("gerenciamento"))

    c = conectar_banco()

    try:
        c.execute("""
            INSERT INTO usuarios(
                usuario, senha, nome, perfil, perfis, ativo, criado_em, senha_temporaria
            )
            VALUES(?,?,?,?,?,?,?,?)
        """, (
            usuario,
            generate_password_hash(senha),
            nome,
            perfil_principal,
            perfis_texto,
            1,
            datetime.now().strftime("%d/%m/%Y %H:%M"),
            0
        ))
        c.commit()
        flash("Usuário criado com sucesso.", "sucesso")
    except sqlite3.IntegrityError:
        c.rollback()
        flash("Esse usuário já existe.", "erro")
    finally:
        c.close()

    return redirect(url_for("gerenciamento"))


@app.route(
    "/alterar-perfil/<int:usuario_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def alterar_perfil(usuario_id):

    perfis_selecionados = [
        p for p in request.form.getlist("perfis")
        if p in PERFIS
    ]

    if not perfis_selecionados:
        perfil_legado = request.form.get("perfil", "operador")
        if perfil_legado in PERFIS:
            perfis_selecionados = [perfil_legado]
        else:
            perfis_selecionados = ["operador"]

    if "admin" in perfis_selecionados:
        perfis_selecionados = ["admin"] + [p for p in perfis_selecionados if p != "admin"]

    perfil_principal = perfis_selecionados[0]
    perfis_texto = ",".join(perfis_selecionados)

    c = conectar_banco()
    usuario = c.execute("SELECT * FROM usuarios WHERE id = ?", (usuario_id,)).fetchone()

    if not usuario:
        c.close()
        flash("Usuário não encontrado.", "erro")
        return redirect(url_for("gerenciamento"))

    tinha_admin = usuario_tem_perfil(usuario, "admin")
    vai_perder_admin = tinha_admin and "admin" not in perfis_selecionados

    if vai_perder_admin:
        total = c.execute("""
            SELECT COUNT(*) total
            FROM usuarios
            WHERE ativo = 1
              AND (perfil = 'admin' OR instr(',' || COALESCE(perfis, '') || ',', ',admin,') > 0)
              AND id != ?
        """, (usuario_id,)).fetchone()["total"]

        if total <= 0:
            c.close()
            flash("Não é possível retirar o cargo do último administrador ativo.", "erro")
            return redirect(url_for("gerenciamento"))

    c.execute("""
        UPDATE usuarios
        SET perfil = ?, perfis = ?
        WHERE id = ?
    """, (perfil_principal, perfis_texto, usuario_id))

    c.commit()
    c.close()

    flash("Cargos atualizados com sucesso.", "sucesso")
    return redirect(url_for("gerenciamento"))


# =========================================================
# ATIVAR / DESATIVAR
# =========================================================

@app.route(
    "/alterar-status-usuario/<int:usuario_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def alterar_status_usuario(
    usuario_id
):

    atual = obter_usuario()

    if atual and atual["id"] == usuario_id:

        flash(
            "Você não pode desativar seu próprio usuário.",
            "erro"
        )

        return redirect(
            url_for("gerenciamento")
        )

    c = conectar_banco()

    usuario = c.execute("""
        SELECT *
        FROM usuarios
        WHERE id = ?
    """, (
        usuario_id,
    )).fetchone()

    if not usuario:

        c.close()

        flash(
            "Usuário não encontrado.",
            "erro"
        )

        return redirect(
            url_for("gerenciamento")
        )

    if (
        usuario_tem_perfil(usuario, "admin")
        and usuario["ativo"] == 1
    ):

        total = c.execute("""
            SELECT COUNT(*) total
            FROM usuarios
            WHERE ativo = 1
              AND (perfil = 'admin' OR instr(',' || COALESCE(perfis, '') || ',', ',admin,') > 0)
        """).fetchone()["total"]

        if total <= 1:

            c.close()

            flash(
                "Não é possível desativar o último administrador ativo.",
                "erro"
            )

            return redirect(
                url_for("gerenciamento")
            )

    novo_status = (
        0
        if usuario["ativo"]
        else 1
    )

    c.execute("""
        UPDATE usuarios
        SET ativo = ?
        WHERE id = ?
    """, (
        novo_status,
        usuario_id
    ))

    c.commit()
    c.close()

    flash(
        "Status do usuário atualizado.",
        "sucesso"
    )

    return redirect(
        url_for("gerenciamento")
    )


# =========================================================
# RESETAR SENHA
# =========================================================

@app.route(
    "/alterar-senha/<int:usuario_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def alterar_senha(usuario_id):

    atual = obter_usuario()

    if atual and atual["id"] == usuario_id:

        flash(
            "Você não pode resetar sua própria senha por esta função.",
            "erro"
        )

        return redirect(
            url_for("gerenciamento")
        )

    c = conectar_banco()

    usuario = c.execute("""
        SELECT *
        FROM usuarios
        WHERE id = ?
    """, (
        usuario_id,
    )).fetchone()

    if not usuario:

        c.close()

        flash(
            "Usuário não encontrado.",
            "erro"
        )

        return redirect(
            url_for("gerenciamento")
        )

    c.execute("""
        UPDATE usuarios
        SET
            senha = ?,
            senha_temporaria = 1
        WHERE id = ?
    """, (
        generate_password_hash("temp"),
        usuario_id
    ))

    c.commit()
    c.close()

    flash(
        f"Senha do usuário {usuario['usuario']} resetada para TEMP.",
        "sucesso"
    )

    return redirect(
        url_for("gerenciamento")
    )


# =========================================================
# EXCLUIR USUÁRIO
# =========================================================

@app.route(
    "/excluir-usuario/<int:usuario_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def excluir_usuario(usuario_id):

    atual = obter_usuario()

    if atual and atual["id"] == usuario_id:

        flash(
            "Você não pode excluir seu próprio usuário.",
            "erro"
        )

        return redirect(
            url_for("gerenciamento")
        )

    c = conectar_banco()

    usuario = c.execute("""
        SELECT *
        FROM usuarios
        WHERE id = ?
    """, (
        usuario_id,
    )).fetchone()

    if not usuario:

        c.close()

        flash(
            "Usuário não encontrado.",
            "erro"
        )

        return redirect(
            url_for("gerenciamento")
        )

    if usuario_tem_perfil(usuario, "admin"):

        total = c.execute("""
            SELECT COUNT(*) total
            FROM usuarios
            WHERE ativo = 1
              AND (perfil = 'admin' OR instr(',' || COALESCE(perfis, '') || ',', ',admin,') > 0)
        """).fetchone()["total"]

        if total <= 1:

            c.close()

            flash(
                "Não é possível excluir o último administrador ativo.",
                "erro"
            )

            return redirect(
                url_for("gerenciamento")
            )

    c.execute("""
        DELETE FROM usuarios
        WHERE id = ?
    """, (
        usuario_id,
    ))

    c.commit()
    c.close()

    flash(
        "Usuário excluído com sucesso.",
        "sucesso"
    )

    return redirect(
        url_for("gerenciamento")
    )


# =========================================================
# EXCLUIR SOLICITAÇÃO / RECEBIMENTO
# SOMENTE ADMIN
# =========================================================

@app.route(
    "/excluir-recebimento/<int:recebimento_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def excluir_recebimento(
    recebimento_id
):

    c = conectar_banco()

    recebimento = c.execute("""
        SELECT *
        FROM recebimentos
        WHERE id = ?
    """, (
        recebimento_id,
    )).fetchone()

    if not recebimento:

        c.close()

        flash(
            "Solicitação não encontrada.",
            "erro"
        )

        return redirect(
            url_for("historico")
        )

    # Apaga o arquivo da NF se existir

    if recebimento["arquivo_nfe"]:

        caminho = os.path.join(
            PASTA_NFES,
            recebimento["arquivo_nfe"]
        )

        try:

            if os.path.isfile(caminho):
                os.remove(caminho)

        except OSError:
            pass

    # Primeiro os itens

    c.execute("""
        DELETE FROM itens
        WHERE recebimento_id = ?
    """, (
        recebimento_id,
    ))

    # Depois o recebimento

    c.execute("""
        DELETE FROM recebimentos
        WHERE id = ?
    """, (
        recebimento_id,
    ))

    c.commit()
    c.close()

    flash(
        "Solicitação excluída com sucesso.",
        "sucesso"
    )

    origem = request.form.get(
        "origem",
        "historico"
    )

    if origem == "pendentes":
        return redirect(
            url_for("pendentes")
        )

    if origem == "conferencia":
        return redirect(
            url_for("conferencia")
        )

    if origem == "notas_fiscais":
        return redirect(
            url_for("notas_fiscais")
        )

    return redirect(
        url_for("historico")
    )


# =========================================================
# EXCLUIR SOMENTE NF-E
# SOMENTE ADMIN
# =========================================================

@app.route(
    "/excluir-nfe/<int:recebimento_id>",
    methods=["POST"]
)
@login_obrigatorio
@requer_permissao("gerenciamento")
def excluir_nfe(recebimento_id):

    c = conectar_banco()

    registro = c.execute("""
        SELECT arquivo_nfe
        FROM recebimentos
        WHERE id = ?
    """, (
        recebimento_id,
    )).fetchone()

    if not registro:

        c.close()

        flash(
            "Registro não encontrado.",
            "erro"
        )

        return redirect(
            url_for("notas_fiscais")
        )

    # Apaga arquivo físico

    if registro["arquivo_nfe"]:

        caminho = os.path.join(
            PASTA_NFES,
            registro["arquivo_nfe"]
        )

        try:

            if os.path.isfile(caminho):
                os.remove(caminho)

        except OSError:
            pass

    # Mantém a solicitação,
    # mas remove os dados da NF

    c.execute("""
        UPDATE recebimentos

        SET
            numero_nfe = NULL,
            responsavel_entrada = NULL,
            data_entrada = NULL,
            arquivo_nfe = NULL,
            status = 'Aguardando entrada',
            conferencia_resultado = 'rejeitado',
            motivo_rejeicao = 'NF-e removida pelo administrador.',
            responsavel_conferencia = NULL,
            data_conferencia = NULL

        WHERE id = ?
    """, (
        recebimento_id,
    ))

    c.commit()
    c.close()

    registrar_evento_fluxo(recebimento_id, "NF-e removida", obter_usuario()["nome"], datetime.now().strftime("%d/%m/%Y %H:%M"), "NF-e removida pelo administrador.")

    flash(
        "NF-e excluída com sucesso. O recebimento permanece conferido e aguarda uma nova NF-e.",
        "sucesso"
    )

    return redirect(
        url_for("notas_fiscais")
    )


# =========================================================
# INICIAR
# =========================================================

if __name__ == "__main__":

    criar_banco()

    print("")
    print("========================================")
    print("       CONTROLE DE CARENAGENS")
    print("========================================")
    print("")
    print("Acesse:")
    print("http://127.0.0.1:5000")
    print("")
    print("Usuário inicial: admin")
    print("Senha inicial: admin123")
    print("")

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False
    )